# -*- coding: utf-8 -*-
"""Génération du classeur Excel des bons de carburant.

Le fichier produit reprend la mise en forme du classeur tenu à la main :
une feuille par mois nommée « Janvier 2026 », les huit mêmes colonnes
(Numéro, Date, Engin, Nom, Carburant, Prix Unitaire, Quantité, Total) et les
lignes triées par date. Il peut donc remplacer le classeur d'origine, ou lui
être réimporté sans perte (l'import relit ce format).

Une feuille de synthèse facultative ajoute les totaux par mois, par carburant,
par engin et par bénéficiaire — ce que le classeur d'origine calculait à la
main en bas de feuille.
"""

import base64
import io
from collections import OrderedDict

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from ..models.secretariat_fuel_voucher import month_label_fr

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - dépendance déclarée au manifeste
    openpyxl = None

_HEADERS = ["Numéro", "Date", "Engin", "Nom", "Carburant",
            "Prix Unitaire", "Quantité", "Total"]
_COLUMN_WIDTHS = [12, 12, 22, 22, 14, 14, 11, 14]
# Caractères interdits par Excel dans un nom de feuille
_FORBIDDEN_SHEET_CHARS = set(r'[]:*?/\\')


class SecretariatFuelExportWizard(models.TransientModel):
    _name = 'secretariat.fuel.export.wizard'
    _description = "Export Excel des bons de carburant"

    date_from = fields.Date(
        string="Du",
        default=lambda self: fields.Date.context_today(self).replace(month=1, day=1),
    )
    date_to = fields.Date(
        string="Au", default=fields.Date.context_today)
    all_dates = fields.Boolean(
        string="Toute la période",
        help="Exporte l'intégralité des bons, sans filtre de date.")

    fuel_type_ids = fields.Many2many(
        'secretariat.fuel.type', string="Carburants",
        help="Laisser vide pour tous les carburants.")
    vehicle_ids = fields.Many2many(
        'secretariat.vehicle', string="Engins",
        help="Laisser vide pour tous les engins.")
    beneficiary_ids = fields.Many2many(
        'secretariat.fuel.beneficiary', string="Bénéficiaires",
        help="Laisser vide pour tous les bénéficiaires.")
    include_cancelled = fields.Boolean(
        string="Inclure les bons annulés", default=False)
    include_summary = fields.Boolean(
        string="Ajouter une feuille de synthèse", default=True)
    include_totals = fields.Boolean(
        string="Ligne de total en bas de feuille", default=True)

    state = fields.Selection(
        [('config', 'Paramètres'), ('done', 'Fichier prêt')],
        default='config', required=True)
    file_data = fields.Binary(string="Classeur", readonly=True, attachment=False)
    file_name = fields.Char(string="Nom du fichier", readonly=True)
    voucher_count = fields.Integer(string="Bons exportés", readonly=True)

    # =========================================================================
    # ACTION
    # =========================================================================

    def action_generate(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_(
                "La bibliothèque Python « openpyxl » est absente du serveur : "
                "la génération Excel est indisponible."))
        if not self.all_dates and self.date_from and self.date_to \
                and self.date_from > self.date_to:
            raise UserError(_("La date de début est postérieure à la date de fin."))

        vouchers = self.env['secretariat.fuel.voucher'].search(
            self._build_domain(), order='date asc, name asc, id asc')
        if not vouchers:
            raise UserError(_(
                "Aucun bon de carburant ne correspond à ces critères."))

        content = self._build_workbook(vouchers)
        self.write({
            'state': 'done',
            'file_data': base64.b64encode(content),
            'file_name': self._build_file_name(),
            'voucher_count': len(vouchers),
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    # =========================================================================
    # SÉLECTION
    # =========================================================================

    def _build_domain(self):
        domain = []
        if not self.all_dates:
            if self.date_from:
                domain.append(('date', '>=', self.date_from))
            if self.date_to:
                domain.append(('date', '<=', self.date_to))
        if not self.include_cancelled:
            domain.append(('state', '!=', 'cancelled'))
        if self.fuel_type_ids:
            domain.append(('fuel_type_id', 'in', self.fuel_type_ids.ids))
        if self.vehicle_ids:
            domain.append(('vehicle_id', 'in', self.vehicle_ids.ids))
        if self.beneficiary_ids:
            domain.append(('beneficiary_id', 'in', self.beneficiary_ids.ids))
        return domain

    def _build_file_name(self):
        if self.all_dates:
            return "bons_carburant_tout.xlsx"
        return "bons_carburant_%s_%s.xlsx" % (self.date_from or '', self.date_to or '')

    # =========================================================================
    # CONSTRUCTION DU CLASSEUR
    # =========================================================================

    def _build_workbook(self, vouchers):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)

        for label, month_vouchers in self._group_by_month(vouchers).items():
            self._write_month_sheet(workbook, label, month_vouchers)

        if self.include_summary:
            self._write_summary_sheet(workbook, vouchers)

        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def _group_by_month(self, vouchers):
        """Regroupe par mois CIVIL de la date du bon.

        Volontairement indépendant du nom des feuilles d'origine : dans le
        classeur repris, la feuille « Janvier 2026 » contenait aussi des bons
        de décembre 2025.

        Le regroupement passe par des listes d'identifiants : réunir des
        recordsets dans une boucle (``months[label] |= voucher``) est
        quadratique et devient sensible dès quelques milliers de bons.
        """
        buckets = OrderedDict()
        for voucher in vouchers:
            buckets.setdefault(month_label_fr(voucher.date), []).append(voucher.id)
        return OrderedDict(
            (label, vouchers.browse(ids)) for label, ids in buckets.items())

    def _write_month_sheet(self, workbook, label, vouchers):
        sheet = workbook.create_sheet(self._safe_sheet_name(workbook, label))
        self._write_header(sheet)

        row_index = 2
        for voucher in vouchers:
            values = [
                self._as_number(voucher.name),
                voucher.date,
                voucher.vehicle_id.name or '',
                voucher.beneficiary_id.name or '',
                voucher.fuel_type_id.name or '',
                voucher.price_unit,
                voucher.quantity,
                voucher.amount_total,
            ]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, col, value)
                if col == 2:
                    cell.number_format = 'DD/MM/YYYY'
                elif col in (6, 8):
                    cell.number_format = '#,##0'
                elif col == 7:
                    cell.number_format = '#,##0.##'
            row_index += 1

        if self.include_totals and vouchers:
            self._write_total_row(sheet, row_index, vouchers)
        self._apply_layout(sheet)
        return sheet

    def _write_header(self, sheet):
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="305496")
        thin = Side(style='thin', color="BFBFBF")
        for col, title in enumerate(_HEADERS, start=1):
            cell = sheet.cell(1, col, title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

    def _write_total_row(self, sheet, row_index, vouchers):
        bold = Font(bold=True)
        top = Border(top=Side(style='thin', color="305496"))
        sheet.cell(row_index, 5, "TOTAL").font = bold
        quantity_cell = sheet.cell(row_index, 7, sum(vouchers.mapped('quantity')))
        amount_cell = sheet.cell(row_index, 8, sum(vouchers.mapped('amount_total')))
        quantity_cell.number_format = '#,##0.##'
        amount_cell.number_format = '#,##0'
        for col in range(1, len(_HEADERS) + 1):
            cell = sheet.cell(row_index, col)
            cell.font = bold
            cell.border = top

    def _write_summary_sheet(self, workbook, vouchers):
        sheet = workbook.create_sheet("Synthèse")
        bold = Font(bold=True)
        row = 1
        sheet.cell(row, 1, "Synthèse des bons de carburant").font = Font(bold=True, size=14)
        row += 2

        for title, key in (
            ("Par mois", lambda v: month_label_fr(v.date)),
            ("Par carburant", lambda v: v.fuel_type_id.name or _("(non précisé)")),
            ("Par engin", lambda v: v.vehicle_id.name or _("(non précisé)")),
            ("Par bénéficiaire", lambda v: v.beneficiary_id.name or _("(non précisé)")),
        ):
            sheet.cell(row, 1, title).font = Font(bold=True, size=12)
            row += 1
            for header_col, header in enumerate(
                    ["Libellé", "Nombre de bons", "Quantité", "Montant"], start=1):
                cell = sheet.cell(row, header_col, header)
                cell.font = bold
                cell.fill = PatternFill("solid", fgColor="D9E1F2")
            row += 1

            buckets = OrderedDict()
            for voucher in vouchers:
                bucket = buckets.setdefault(key(voucher), [0, 0.0, 0.0])
                bucket[0] += 1
                bucket[1] += voucher.quantity
                bucket[2] += voucher.amount_total
            for label, (count, quantity, amount) in sorted(
                    buckets.items(), key=lambda item: -item[1][2]):
                sheet.cell(row, 1, label)
                sheet.cell(row, 2, count)
                sheet.cell(row, 3, quantity).number_format = '#,##0.##'
                sheet.cell(row, 4, amount).number_format = '#,##0'
                row += 1
            sheet.cell(row, 1, "Total").font = bold
            sheet.cell(row, 2, len(vouchers)).font = bold
            total_quantity = sheet.cell(row, 3, sum(vouchers.mapped('quantity')))
            total_amount = sheet.cell(row, 4, sum(vouchers.mapped('amount_total')))
            total_quantity.font = bold
            total_quantity.number_format = '#,##0.##'
            total_amount.font = bold
            total_amount.number_format = '#,##0'
            row += 3

        for col, width in enumerate([32, 16, 14, 16], start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        return sheet

    # =========================================================================
    # HELPERS
    # =========================================================================

    @staticmethod
    def _apply_layout(sheet):
        for col, width in enumerate(_COLUMN_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = "A1:%s1" % get_column_letter(len(_HEADERS))

    @staticmethod
    def _as_number(value):
        """Le classeur d'origine stocke le numéro de bon en nombre."""
        try:
            return int(str(value).strip())
        except (TypeError, ValueError):
            return value

    @staticmethod
    def _safe_sheet_name(workbook, label):
        """Nom de feuille accepté par Excel : 31 caractères, sans []:*?/\\."""
        name = "".join(c for c in label if c not in _FORBIDDEN_SHEET_CHARS)[:31]
        if name not in workbook.sheetnames:
            return name
        suffix = 2
        while "%s (%s)" % (name[:27], suffix) in workbook.sheetnames:
            suffix += 1
        return "%s (%s)" % (name[:27], suffix)

    @api.onchange('all_dates')
    def _onchange_all_dates(self):
        if self.all_dates:
            self.date_from = False
            self.date_to = False
