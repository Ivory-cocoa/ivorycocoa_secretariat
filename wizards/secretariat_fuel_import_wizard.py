# -*- coding: utf-8 -*-
"""Import du classeur Excel des bons de carburant.

Le classeur repris (« yasmina_petroleum.xlsx ») est tenu à la main : une
feuille par période, des lignes de total en bas de feuille, des orthographes
flottantes (« Gasoil » / « gasoil », « Gaz » / « GAZ » / « gaz »), quelques
dates saisies en texte, une cellule « #VALUE! », et surtout des feuilles
recopiées d'un mois sur l'autre — 165 des 1 377 lignes du classeur analysé
sont des doublons stricts.

L'import est donc conçu pour être tolérant ET idempotent :

* détection de la ligne d'en-tête et des colonnes par leur libellé ;
* lignes sans numéro de bon ignorées (ce sont les totaux de bas de feuille) ;
* dates au séparateur oublié (« 25/022026 ») reconstituées, avec une note sur
  le bon créé ;
* prix illisible remplacé par le prix courant du carburant, également signalé ;
* dédoublonnage dans le fichier ET vis-à-vis des bons déjà en base ;
* aucune ligne n'est perdue silencieusement : les lignes non reprises
  ressortent dans un classeur au format même du modèle d'import, cellules
  fautives surlignées, prêt à être corrigé et renvoyé tel quel.

Sur le classeur de référence (1 377 lignes, 15 feuilles) : 1 376 lignes
exploitables, 165 doublons écartés, 1 211 bons créés, 1 ligne rejetée (date
réellement absente).
"""

import base64
import io
import logging
from datetime import date, datetime, timedelta

from markupsafe import Markup, escape

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from ..models.secretariat_fuel_type import normalize_label

_logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - dépendance déclarée au manifeste
    openpyxl = None

# Correspondance libellé de colonne (normalisé) -> clé interne
_COLUMN_ALIASES = {
    'numero': 'name',
    'num': 'name',
    'n bon': 'name',
    'numero du bon': 'name',
    'date': 'date',
    'engin': 'vehicle',
    'vehicule': 'vehicle',
    'nom': 'beneficiary',
    'beneficiaire': 'beneficiary',
    'carburant': 'fuel',
    'produit': 'fuel',
    'prix unitaire': 'price',
    'prix': 'price',
    'pu': 'price',
    'quantite': 'quantity',
    'qte': 'quantity',
    'total': 'total',
    'montant': 'total',
}
_REQUIRED_COLUMNS = ('name', 'date', 'fuel', 'quantity')

# Origine des numéros de série de dates Excel (système 1900, décalage inclus)
_EXCEL_EPOCH = datetime(1899, 12, 30)
# Taille maximale du classeur accepté (le classeur de référence pèse ~200 Ko)
_MAX_FILE_SIZE = 25 * 1024 * 1024
# Écart toléré entre le total du fichier et prix × quantité
_TOTAL_TOLERANCE = 1.0
# Nombre de lignes scannées à la recherche de l'en-tête
_HEADER_SCAN_ROWS = 10

# --- Modèle vierge proposé au téléchargement ------------------------------
_TEMPLATE_HEADERS = ["Numéro", "Date", "Engin", "Nom", "Carburant",
                     "Prix Unitaire", "Quantité", "Total"]
_TEMPLATE_WIDTHS = [12, 12, 24, 24, 16, 14, 11, 14]
# Lignes préformatées : au-delà, la saisie continue sans mise en forme, et
# l'import lit la feuille jusqu'à sa dernière ligne remplie.
_TEMPLATE_ROWS = 300

_TEMPLATE_GUIDE = [
    ("Numéro",
     "Le numéro imprimé sur la souche. Obligatoire — une ligne sans numéro "
     "est ignorée (c'est ainsi que les lignes de total en bas de feuille sont "
     "écartées). Le même numéro peut revenir plusieurs fois : un bon Super et "
     "un bon Lubrifiant peuvent porter le même numéro le même jour."),
    ("Date",
     "La date du bon. Obligatoire. Format jj/mm/aaaa, ou une vraie date "
     "Excel."),
    ("Engin",
     "Le véhicule, la moto, l'engin d'usine… tel qu'écrit sur le bon. "
     "Facultatif : un bon sans engin est repris et marqué « à compléter »."),
    ("Nom",
     "Le bénéficiaire : un employé, un service (Usine, Personnel) ou un "
     "organisme (Douanes, CCC). Facultatif, même remarque."),
    ("Carburant",
     "Super, Gasoil, Lubrifiant, Gaz… Obligatoire. Choisissez-le dans la "
     "liste déroulante ; les accents et la casse n'ont pas d'importance."),
    ("Prix Unitaire",
     "Le prix appliqué à ce bon. S'il est laissé vide, le prix courant du "
     "carburant est repris automatiquement, et le bon en porte la mention."),
    ("Quantité",
     "Obligatoire et strictement positive. En litres pour les carburants, à "
     "l'unité pour les lubrifiants et le gaz."),
    ("Total",
     "Calculé tout seul. Vous pouvez le laisser tel quel : à l'import, c'est "
     "prix × quantité qui fait foi."),
]

_TEMPLATE_NOTES = [
    "Ne changez ni le nom ni l'ordre des colonnes : c'est à leur libellé que "
    "le fichier est reconnu. En revanche, vous pouvez insérer des lignes "
    "au-dessus de l'en-tête (un titre, un logo) : il est cherché dans les dix "
    "premières lignes.",
    "Vous pouvez créer plusieurs feuilles — une par mois, par exemple. "
    "L'assistant les proposera toutes, et vous choisirez celles à reprendre.",
    "Réimporter un fichier déjà repris ne crée aucun doublon : un bon est "
    "reconnu à son numéro, sa date, son engin, son carburant et sa quantité. "
    "Vous pouvez donc corriger le fichier et le renvoyer entier sans crainte.",
    "Les engins, bénéficiaires et carburants absents de la base sont créés "
    "automatiquement — décochez l'option si vous préférez les saisir vous-même "
    "avant l'import.",
]


class _LabelResolver:
    """Donne à chaque libellé du fichier un identifiant stable.

    La clé de dédoublonnage d'un bon contient l'identifiant de l'engin et du
    carburant. Il faut donc pouvoir la calculer AVANT toute création de
    référentiel — sinon l'import crée des engins pour des lignes qu'il écarte
    ensuite comme doublons.

    * un libellé connu prend l'identifiant réel du référentiel : la clé se
      compare alors directement à celles des bons déjà en base ;
    * un libellé inconnu prend un identifiant négatif, unique et stable : deux
      engins inconnus DIFFÉRENTS ne sont pas confondus, et aucun bon existant
      ne peut porter cet identifiant ;
    * un libellé vide vaut 0, comme un engin absent en base.
    """

    def __init__(self, known_ids):
        self._known = known_ids
        self._unknown = {}
        self._next = -1

    def key_id(self, label):
        normalized = normalize_label(label)
        if not normalized:
            return 0
        if normalized in self._known:
            return self._known[normalized]
        if normalized not in self._unknown:
            self._unknown[normalized] = self._next
            self._next -= 1
        return self._unknown[normalized]


class SecretariatFuelImportSheet(models.TransientModel):
    """Une feuille du classeur, telle que proposée à l'import."""

    _name = 'secretariat.fuel.import.sheet'
    _description = "Feuille du classeur à importer"
    _order = 'sequence, id'

    wizard_id = fields.Many2one(
        'secretariat.fuel.import.wizard', required=True, ondelete='cascade')
    sequence = fields.Integer(string="Ordre")
    name = fields.Char(string="Feuille", readonly=True)
    selected = fields.Boolean(string="Importer", default=True)
    row_count = fields.Integer(string="Lignes de bon", readonly=True)
    new_count = fields.Integer(
        string="Nouvelles", readonly=True,
        help="Lignes qui seront créées : ni déjà en base, ni en doublon dans "
             "le fichier.")
    duplicate_count = fields.Integer(
        string="Doublons", readonly=True,
        help="Lignes déjà présentes en base ou déjà vues dans une feuille "
             "précédente du même fichier.")
    error_count = fields.Integer(string="En erreur", readonly=True)


class SecretariatFuelImportWizard(models.TransientModel):
    _name = 'secretariat.fuel.import.wizard'
    _description = "Import des bons de carburant depuis Excel"

    state = fields.Selection(
        [('upload', 'Fichier'), ('analyzed', 'Analyse'), ('done', 'Résultat')],
        default='upload',
        required=True,
    )

    file_data = fields.Binary(string="Fichier Excel (.xlsx)", attachment=False)
    file_name = fields.Char(string="Nom du fichier")

    sheet_ids = fields.One2many(
        'secretariat.fuel.import.sheet', 'wizard_id', string="Feuilles")
    preview_html = fields.Html(
        string="Aperçu", readonly=True, sanitize=False,
        help="Les premières lignes du fichier, telles qu'elles ont été "
             "comprises. C'est le moment de vérifier que les colonnes sont "
             "bien tombées en face.")

    create_missing = fields.Boolean(
        string="Créer les référentiels manquants",
        default=True,
        help="Crée automatiquement les engins, bénéficiaires et types de "
             "carburant rencontrés dans le fichier et absents de la base.",
    )
    skip_duplicates = fields.Boolean(
        string="Ignorer les doublons",
        default=True,
        help="Un bon est considéré comme déjà importé s'il a le même numéro, "
             "la même date, le même engin, le même carburant et la même "
             "quantité. À décocher uniquement pour forcer un ré-import.",
    )
    target_state = fields.Selection(
        [('draft', 'Brouillon'), ('confirmed', 'Validé')],
        string="État des bons importés",
        default='confirmed',
        required=True,
        help="Les bons repris de l'ancien classeur ont déjà été honorés : "
             "« Validé » est l'état attendu.",
    )
    supplier_id = fields.Many2one(
        'res.partner', string="Station",
        domain="[('is_company', '=', True)]",
        help="Station-service appliquée à tous les bons importés (facultatif).")

    # --- Résultat ---
    imported_count = fields.Integer(string="Bons créés", readonly=True)
    duplicate_count = fields.Integer(string="Doublons ignorés", readonly=True)
    error_count = fields.Integer(string="Lignes en erreur", readonly=True)
    created_vehicle_count = fields.Integer(string="Engins créés", readonly=True)
    created_beneficiary_count = fields.Integer(string="Bénéficiaires créés", readonly=True)
    created_fuel_count = fields.Integer(string="Carburants créés", readonly=True)
    result_html = fields.Html(string="Compte rendu", readonly=True, sanitize=False)
    error_file = fields.Binary(string="Lignes à corriger (Excel)", readonly=True)
    error_file_name = fields.Char(readonly=True)

    # --- Modèle vierge ---
    template_file = fields.Binary(readonly=True, attachment=False)
    template_file_name = fields.Char(
        readonly=True, default="modele_bons_carburant.xlsx")

    # =========================================================================
    # ÉTAPE 0 — LE MODÈLE VIERGE
    # =========================================================================

    def action_download_template(self):
        """Produit un classeur vierge aux bonnes colonnes, et le télécharge.

        Sans lui, la secrétaire doit deviner les libellés exacts attendus —
        c'est la première marche de l'import, et la plus haute. Le fichier est
        déposé dans un champ binaire de l'assistant lui-même : pas de pièce
        jointe qui traîne, l'enregistrement transitoire est purgé tout seul.
        """
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_(
                "La bibliothèque Python « openpyxl » est absente du serveur : "
                "la génération du modèle est indisponible."))
        self.write({
            'template_file': base64.b64encode(self._build_template()),
            'template_file_name': "modele_bons_carburant.xlsx",
        })
        return {
            'type': 'ir.actions.act_url',
            'target': 'self',
            'url': '/web/content?model=%s&id=%s&field=template_file'
                   '&filename_field=template_file_name&download=true' % (
                       self._name, self.id),
        }

    def _build_template(self):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        self._write_template_sheet(workbook)
        self._write_template_guide(workbook)
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def _write_template_sheet(self, workbook):
        """La feuille à remplir : mêmes colonnes que le carnet."""
        sheet = workbook.create_sheet("Bons de carburant")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="305496")
        thin = Side(style='thin', color="BFBFBF")
        for column, title in enumerate(_TEMPLATE_HEADERS, start=1):
            cell = sheet.cell(1, column, title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for column, width in enumerate(_TEMPLATE_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = 'A2'

        # Formats et formule de total sur les premières lignes : la secrétaire
        # voit le montant se calculer pendant qu'elle saisit. À l'import, c'est
        # de toute façon prix × quantité qui fait foi.
        for row in range(2, _TEMPLATE_ROWS + 2):
            sheet.cell(row, 2).number_format = 'DD/MM/YYYY'
            sheet.cell(row, 6).number_format = '#,##0'
            sheet.cell(row, 7).number_format = '#,##0.##'
            total = sheet.cell(row, 8)
            total.value = '=IF(OR(F%(r)s="",G%(r)s=""),"",F%(r)s*G%(r)s)' % {'r': row}
            total.number_format = '#,##0'

        # Liste déroulante des carburants connus, pour éviter les fautes de
        # frappe. Excel plafonne la formule à 255 caractères : au-delà, on
        # laisse la colonne libre (l'import rapproche les orthographes).
        names = self.env['secretariat.fuel.type'].search([]).mapped('name')
        joined = ",".join(name.replace('"', '') for name in names)
        if names and len(joined) <= 250:
            validation = DataValidation(
                type='list', formula1='"%s"' % joined, allow_blank=True)
            validation.error = "Choisissez un carburant de la liste."
            validation.errorTitle = "Carburant inconnu"
            sheet.add_data_validation(validation)
            validation.add('E2:E%s' % (_TEMPLATE_ROWS + 1))
        return sheet

    def _write_template_guide(self, workbook):
        """Mode d'emploi.

        Disposé en deux colonnes à dessein : la détection d'en-tête exige de
        trouver Numéro, Date, Carburant ET Quantité sur UNE MÊME ligne. Ici
        chaque libellé est seul sur sa ligne — cette feuille ne peut donc pas
        être prise pour une feuille de données.
        """
        sheet = workbook.create_sheet("Mode d'emploi")
        sheet.column_dimensions['A'].width = 20
        sheet.column_dimensions['B'].width = 96

        sheet.cell(1, 1, "Comment remplir ce classeur").font = Font(bold=True, size=14)
        row = 3
        for label, explanation in _TEMPLATE_GUIDE:
            cell = sheet.cell(row, 1, label)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical='top')
            body = sheet.cell(row, 2, explanation)
            body.alignment = Alignment(vertical='top', wrap_text=True)
            row += 1

        row += 1
        sheet.cell(row, 1, "Bon à savoir").font = Font(bold=True, size=12)
        row += 1
        for note in _TEMPLATE_NOTES:
            sheet.cell(row, 2, note).alignment = Alignment(
                vertical='top', wrap_text=True)
            row += 1
        return sheet

    # =========================================================================
    # ÉTAPE 1 — ANALYSE
    # =========================================================================

    def action_analyze(self):
        """Ouvre le classeur et propose ses feuilles, avec un décompte.

        Le décompte est calculé exactement comme le fera l'import : même
        lecture, même clé de dédoublonnage. Ce que l'écran annonce est donc ce
        que l'import produira.
        """
        self.ensure_one()
        workbook = self._load_workbook()

        parsed = []
        for index, sheet in enumerate(workbook.worksheets):
            rows, errors = self._parse_sheet(sheet)
            parsed.append((index, sheet.title, rows, errors))

        all_rows = [row for _i, _t, rows, _e in parsed for row in rows]
        existing = self._existing_keys(self._date_bounds(all_rows))
        resolvers = self._label_resolvers()
        seen = set()

        lines = [(5, 0, 0)]
        for index, title, rows, errors in parsed:
            new_count = duplicate_count = 0
            for row in rows:
                key = self._row_key(row, resolvers)
                if key in seen or key in existing:
                    duplicate_count += 1
                else:
                    seen.add(key)
                    new_count += 1
            lines.append((0, 0, {
                'sequence': index,
                'name': title,
                'selected': True,
                'row_count': len(rows),
                'new_count': new_count,
                'duplicate_count': duplicate_count,
                'error_count': len(errors),
            }))
        self.write({
            'sheet_ids': lines,
            'state': 'analyzed',
            'preview_html': self._build_preview(all_rows),
        })
        return self._reopen()

    def action_back(self):
        """Retour à l'étape « Fichier » sans fermer l'assistant."""
        self.ensure_one()
        self.write({'state': 'upload', 'sheet_ids': [(5, 0, 0)]})
        return self._reopen()

    # =========================================================================
    # ÉTAPE 2 — IMPORT
    # =========================================================================

    def action_import(self):
        self.ensure_one()
        selected = self.sheet_ids.filtered('selected').mapped('name')
        if not selected:
            raise UserError(_("Sélectionnez au moins une feuille à importer."))

        workbook = self._load_workbook()
        Voucher = self.env['secretariat.fuel.voucher']

        # 1. Lecture des feuilles retenues.
        rows, errors = [], []
        for sheet in workbook.worksheets:
            if sheet.title not in selected:
                continue
            sheet_rows, sheet_errors = self._parse_sheet(sheet)
            rows.extend(sheet_rows)
            errors.extend(sheet_errors)

        # 2. Dédoublonnage AVANT toute création de référentiel : une ligne
        #    écartée ne doit pas laisser derrière elle un engin fantôme.
        existing = self._existing_keys(self._date_bounds(rows)) \
            if self.skip_duplicates else set()
        resolvers = self._label_resolvers()
        seen = set()
        kept = []
        duplicates = 0
        for row in rows:
            key = self._row_key(row, resolvers)
            if self.skip_duplicates and (key in seen or key in existing):
                duplicates += 1
                continue
            seen.add(key)
            kept.append(row)

        # 3. Résolution (et création) des référentiels, sur les seules lignes
        #    effectivement reprises.
        referentials = {'vehicle': dict(), 'beneficiary': dict(), 'fuel': dict()}
        created_counts = {'vehicle': 0, 'beneficiary': 0, 'fuel': 0}
        values_list = []
        for row in kept:
            try:
                values_list.append(self._prepare_voucher_values(
                    row, referentials, created_counts))
            except UserError as exc:
                errors.append(dict(row, error=str(exc), culprit='fuel'))

        vouchers = Voucher.create(values_list) if values_list else Voucher

        self.write({
            'state': 'done',
            'imported_count': len(vouchers),
            'duplicate_count': duplicates,
            'error_count': len(errors),
            'created_vehicle_count': created_counts['vehicle'],
            'created_beneficiary_count': created_counts['beneficiary'],
            'created_fuel_count': created_counts['fuel'],
            'result_html': self._build_report(len(vouchers), duplicates, errors),
            'error_file': self._build_error_file(errors),
            'error_file_name': errors and "bons_carburant_a_corriger.xlsx" or False,
        })
        _logger.info(
            "Import bons de carburant : %s créés, %s doublons, %s erreurs (%s)",
            len(vouchers), duplicates, len(errors), self.file_name)
        return self._reopen()

    def action_view_imported(self):
        """Ouvre la liste des bons importés depuis ce fichier."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _("Bons importés"),
            'res_model': 'secretariat.fuel.voucher',
            'view_mode': 'tree,form,pivot,graph',
            'domain': [('imported', '=', True)],
        }

    # =========================================================================
    # LECTURE DU CLASSEUR
    # =========================================================================

    def _load_workbook(self):
        if openpyxl is None:
            raise UserError(_(
                "La bibliothèque Python « openpyxl » est absente du serveur : "
                "l'import Excel est indisponible."))
        if not self.file_data:
            raise UserError(_("Choisissez d'abord un fichier Excel (.xlsx)."))
        if self.file_name and not self.file_name.lower().endswith(('.xlsx', '.xlsm')):
            raise UserError(_(
                "Format non pris en charge : %s. Enregistrez le classeur au "
                "format .xlsx puis relancez l'import.", self.file_name))
        try:
            content = base64.b64decode(self.file_data)
        except Exception:
            raise UserError(_("Le fichier envoyé est illisible. Renvoyez-le."))
        if len(content) > _MAX_FILE_SIZE:
            raise UserError(_(
                "Le classeur dépasse %s Mo. Découpez-le par période et "
                "importez les morceaux l'un après l'autre.",
                _MAX_FILE_SIZE // (1024 * 1024)))
        try:
            return openpyxl.load_workbook(
                io.BytesIO(content), data_only=True, read_only=False)
        except Exception as exc:
            raise UserError(_(
                "Le fichier n'a pas pu être ouvert : %s\n\nVérifiez qu'il "
                "s'agit bien d'un classeur Excel (.xlsx) non protégé.", exc))

    def _parse_sheet(self, sheet):
        """Retourne (lignes exploitables, lignes en erreur) d'une feuille."""
        columns = self._detect_columns(sheet)
        if not columns:
            return [], []

        header_row, mapping = columns
        rows, errors = [], []
        for row_index in range(header_row + 1, sheet.max_row + 1):
            raw = {
                key: sheet.cell(row_index, col).value
                for key, col in mapping.items()
            }
            # Ligne sans numéro : total de bas de feuille ou ligne vide.
            if raw.get('name') in (None, '') or not str(raw['name']).strip():
                continue
            row = {
                'sheet': sheet.title,
                'row': row_index,
                'raw_name': raw.get('name'),
                'raw_date': raw.get('date'),
                'vehicle': self._clean_text(raw.get('vehicle')),
                'beneficiary': self._clean_text(raw.get('beneficiary')),
                'fuel': self._clean_text(raw.get('fuel')),
            }
            row['name'] = self._clean_number_as_text(raw.get('name'))
            row['date'] = self._to_date(raw.get('date'))
            row['date_repaired'] = False
            if row['date'] is None and raw.get('date'):
                row['date'] = self._repair_date(raw['date'])
                row['date_repaired'] = bool(row['date'])
            row['price'] = self._to_float(raw.get('price'))
            row['quantity'] = self._to_float(raw.get('quantity'))
            row['total'] = self._to_float(raw.get('total'))

            # « culprit » désigne la colonne fautive : c'est elle que le
            # classeur des lignes à corriger surligne en rouge.
            if row['date'] is None:
                errors.append(dict(
                    row, error=_("Date illisible ou absente"), culprit='date'))
                continue
            if not row['fuel']:
                errors.append(dict(
                    row, error=_("Carburant absent"), culprit='fuel'))
                continue
            if row['quantity'] is None or row['quantity'] <= 0:
                errors.append(dict(
                    row, error=_("Quantité absente ou nulle"), culprit='quantity'))
                continue
            rows.append(row)
        return rows, errors

    def _detect_columns(self, sheet):
        """Localise la ligne d'en-tête et associe chaque colonne à une clé."""
        max_col = min(sheet.max_column or 0, 40)
        for row_index in range(1, min(_HEADER_SCAN_ROWS, sheet.max_row or 0) + 1):
            mapping = {}
            for col in range(1, max_col + 1):
                key = _COLUMN_ALIASES.get(normalize_label(sheet.cell(row_index, col).value))
                if key and key not in mapping:
                    mapping[key] = col
            if all(required in mapping for required in _REQUIRED_COLUMNS):
                return row_index, mapping
        return None

    # =========================================================================
    # CONVERSIONS TOLÉRANTES
    # =========================================================================

    @staticmethod
    def _clean_text(value):
        if value is None:
            return ''
        return str(value).strip()

    @staticmethod
    def _clean_number_as_text(value):
        """Le numéro de bon est un entier dans le fichier — pas « 7138.0 »."""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @staticmethod
    def _to_float(value):
        if value is None or value == '':
            return None
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
        text = str(value).strip().replace(' ', '').replace(' ', '')
        text = text.replace(',', '.')
        try:
            return float(text)
        except ValueError:
            return None

    @staticmethod
    def _to_date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            # Numéro de série Excel
            try:
                return (_EXCEL_EPOCH + timedelta(days=float(value))).date()
            except (OverflowError, ValueError):
                return None
        if not value:
            return None
        text = str(value).strip()
        for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y',
                    '%d.%m.%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _repair_date(value):
        """Rattrape les dates saisies avec un séparateur manquant.

        Le classeur repris en contient une poignée : « 25/022026 »,
        « 13/03/026 », « 2004/2026 », « 1506/2026 », « 30/072026 ». On ne
        garde que les chiffres et on relit en jour-mois-année ; le bon
        importé porte alors une note signalant la reconstitution.
        """
        digits = ''.join(c for c in str(value) if c.isdigit())
        if len(digits) == 8:
            day, month, year = digits[:2], digits[2:4], int(digits[4:])
        elif len(digits) == 7:
            # Année sur 3 chiffres : « 13/03/026 »
            day, month, year = digits[:2], digits[2:4], 2000 + int(digits[4:])
        elif len(digits) == 6:
            day, month, year = digits[:2], digits[2:4], 2000 + int(digits[4:])
        else:
            return None
        try:
            return date(year, int(month), int(day))
        except ValueError:
            # Dernière tentative : mois et jour intervertis
            try:
                return date(year, int(day), int(month))
            except ValueError:
                return None

    # =========================================================================
    # PRÉPARATION DES ENREGISTREMENTS
    # =========================================================================

    def _prepare_voucher_values(self, row, referentials, created_counts):
        fuel_type = self._resolve(
            'fuel', row['fuel'], referentials, created_counts,
            self.env['secretariat.fuel.type'])
        if not fuel_type:
            raise UserError(_(
                "Carburant « %s » inconnu (création des référentiels "
                "désactivée)", row['fuel']))

        vehicle = self._resolve(
            'vehicle', row['vehicle'], referentials, created_counts,
            self.env['secretariat.vehicle']) if row['vehicle'] else None
        beneficiary = self._resolve(
            'beneficiary', row['beneficiary'], referentials, created_counts,
            self.env['secretariat.fuel.beneficiary']) if row['beneficiary'] else None

        price = row['price']
        notes = []
        if row.get('date_repaired'):
            notes.append(_(
                "Date reconstituée à partir de la saisie « %(raw)s » du "
                "fichier — à vérifier.", raw=row['raw_date']))
        if price is None or price <= 0:
            price = fuel_type.price
            notes.append(_(
                "Prix unitaire absent du fichier — prix courant du carburant "
                "appliqué (%s).", price))
        computed_total = price * row['quantity']
        if row['total'] is not None and abs(row['total'] - computed_total) > _TOTAL_TOLERANCE:
            notes.append(_(
                "Total du fichier (%(file)s) différent de prix × quantité "
                "(%(computed)s) — c'est le calcul qui fait foi.",
                file=row['total'], computed=computed_total))

        return {
            'name': row['name'],
            'date': row['date'],
            'vehicle_id': vehicle.id if vehicle else False,
            'beneficiary_id': beneficiary.id if beneficiary else False,
            'fuel_type_id': fuel_type.id,
            'price_unit': price,
            'quantity': row['quantity'],
            'state': self.target_state,
            'supplier_id': self.supplier_id.id or False,
            'imported': True,
            'import_source': "%s — %s (ligne %s)" % (
                self.file_name or _("classeur"), row['sheet'], row['row']),
            'note': "\n".join(notes) or False,
        }

    def _resolve(self, kind, label, referentials, created_counts, model):
        """Retrouve (ou crée) un référentiel, avec un cache par libellé."""
        key = normalize_label(label)
        if not key:
            return None
        cache = referentials[kind]
        if key in cache:
            return cache[key]
        record = model.with_context(active_test=False).search(
            [('normalized_name', '=', key)], limit=1)
        if not record and self.create_missing:
            record = model.find_or_create(label)
            created_counts[kind] += 1
        cache[key] = record or None
        return cache[key]

    # =========================================================================
    # HELPERS
    # =========================================================================

    @staticmethod
    def _date_bounds(rows):
        """(date la plus ancienne, date la plus récente) des lignes lues."""
        dates = [row['date'] for row in rows if row.get('date')]
        return (min(dates), max(dates)) if dates else None

    def _existing_keys(self, bounds=None):
        """Clés de dédoublonnage des bons déjà en base.

        Bornées à la plage de dates du fichier : inutile de charger dix ans de
        registre pour importer une feuille de mars.
        """
        domain = []
        if bounds:
            domain = [('date', '>=', bounds[0]), ('date', '<=', bounds[1])]
        keys = set()
        Voucher = self.env['secretariat.fuel.voucher']
        for record in Voucher.search_read(
                domain, ['name', 'date', 'vehicle_id', 'fuel_type_id', 'quantity']):
            keys.add(Voucher._dedup_key(
                record['name'], record['date'],
                record['vehicle_id'] and record['vehicle_id'][0],
                record['fuel_type_id'] and record['fuel_type_id'][0],
                record['quantity']))
        return keys

    def _label_resolvers(self):
        """Un résolveur de libellés par référentiel entrant dans la clé."""
        resolvers = {}
        for kind, model_name in (('vehicle', 'secretariat.vehicle'),
                                 ('fuel', 'secretariat.fuel.type')):
            records = self.env[model_name].with_context(active_test=False).search_read(
                [], ['normalized_name'])
            resolvers[kind] = _LabelResolver(
                {record['normalized_name']: record['id'] for record in records})
        return resolvers

    def _row_key(self, row, resolvers):
        """Clé de dédoublonnage d'une ligne du fichier, avant création."""
        return self.env['secretariat.fuel.voucher']._dedup_key(
            row['name'],
            row['date'],
            resolvers['vehicle'].key_id(row['vehicle']),
            resolvers['fuel'].key_id(row['fuel']),
            row['quantity'],
        )

    def _build_preview(self, rows, limit=8):
        """Les premières lignes, telles qu'elles ont été comprises.

        Les compteurs de l'étape 2 disent combien de lignes seront reprises,
        pas ce qu'elles contiennent. Si une colonne avait été mal reconnue —
        un fichier dont l'en-tête diffère, des colonnes interverties — la
        secrétaire ne s'en apercevait qu'une fois les bons créés. Voir huit
        lignes en face de leur libellé lève le doute avant de s'engager.

        Les libellés absents de la base sont signalés « nouveau » : c'est là
        qu'on repère une faute de frappe qui allait créer un engin de plus.
        """
        if not rows:
            return False

        known = {
            kind: {
                record['normalized_name']
                for record in self.env[model].with_context(
                    active_test=False).search_read([], ['normalized_name'])
            }
            for kind, model in (('vehicle', 'secretariat.vehicle'),
                                ('beneficiary', 'secretariat.fuel.beneficiary'),
                                ('fuel', 'secretariat.fuel.type'))
        }
        prices = {
            fuel.normalized_name: fuel.price
            for fuel in self.env['secretariat.fuel.type'].search([])
        }

        def label(value, kind):
            """Le libellé, suivi d'une pastille quand il est inconnu."""
            if not value:
                return Markup('<span class="text-muted">—</span>')
            normalized = normalize_label(value)
            if normalized in known[kind]:
                return escape(value)
            return Markup('%s <span class="badge text-bg-info">nouveau</span>') \
                % escape(value)

        parts = [Markup(
            '<table class="table table-sm table-striped mb-1">'
            '<thead><tr>'
            '<th>Feuille</th><th>Ligne</th><th>N°</th><th>Date</th>'
            '<th>Engin</th><th>Bénéficiaire</th><th>Carburant</th>'
            '<th class="text-end">Prix</th><th class="text-end">Quantité</th>'
            '<th class="text-end">Total</th>'
            '</tr></thead><tbody>')]
        for row in rows[:limit]:
            price = row['price']
            if price is None or price <= 0:
                price = prices.get(normalize_label(row['fuel']), 0.0)
            parts.append(Markup(
                '<tr><td>%s</td><td>%s</td><td><strong>%s</strong></td><td>%s</td>'
                '<td>%s</td><td>%s</td><td>%s</td>'
                '<td class="text-end">%s</td><td class="text-end">%s</td>'
                '<td class="text-end">%s</td></tr>') % (
                escape(row['sheet']), row['row'], escape(row['name']),
                row['date'].strftime('%d/%m/%Y') if row['date'] else '',
                label(row['vehicle'], 'vehicle'),
                label(row['beneficiary'], 'beneficiary'),
                label(row['fuel'], 'fuel'),
                round(price), round(row['quantity'], 2),
                round(price * row['quantity'])))
        parts.append(Markup('</tbody></table>'))
        if len(rows) > limit:
            parts.append(Markup(
                '<p class="text-muted mb-0"><em>… et %s autre(s) ligne(s).</em></p>')
                % (len(rows) - limit))
        return Markup('').join(parts)

    def _build_report(self, imported, duplicates, errors):
        """Compte rendu HTML.

        Tout ce qui vient du classeur (nom de feuille, message d'erreur) est
        échappé : le champ est déclaré ``sanitize=False`` pour préserver la
        mise en forme, il ne doit donc rien recevoir de brut.
        """
        parts = [Markup('<div class="o_secretariat_import_report">')]
        parts.append(Markup('<p><strong>%s</strong> bon(s) créé(s).</p>') % imported)
        if duplicates:
            parts.append(Markup(
                '<p>%s ligne(s) ignorée(s) car déjà présentes en base ou en '
                'double dans le fichier.</p>') % duplicates)
        created = []
        if self.created_vehicle_count:
            created.append(_("%s engin(s)", self.created_vehicle_count))
        if self.created_beneficiary_count:
            created.append(_("%s bénéficiaire(s)", self.created_beneficiary_count))
        if self.created_fuel_count:
            created.append(_("%s type(s) de carburant", self.created_fuel_count))
        if created:
            parts.append(Markup('<p>Référentiels créés : %s.</p>')
                         % escape(", ".join(created)))
        if errors:
            parts.append(Markup(
                '<p class="text-danger"><strong>%s ligne(s) rejetée(s)</strong> — '
                'ouvrez le classeur ci-dessous, corrigez les cellules en rouge '
                "et renvoyez-le tel quel à l'assistant : les lignes déjà reprises "
                'ne seront pas dupliquées.</p>')
                % len(errors))
            parts.append(Markup('<ul>'))
            for error in errors[:10]:
                parts.append(Markup('<li>%s, ligne %s : %s</li>') % (
                    escape(error.get('sheet') or ''),
                    error.get('row') or '',
                    escape(error.get('error') or '')))
            parts.append(Markup('</ul>'))
            if len(errors) > 10:
                parts.append(Markup('<p><em>… et %s autre(s).</em></p>')
                             % (len(errors) - 10))
        parts.append(Markup('</div>'))
        return Markup('').join(parts)

    def _build_error_file(self, errors):
        """Les lignes non reprises, dans le format même du modèle d'import.

        C'est le point : la secrétaire corrige les cellules surlignées et
        renvoie CE fichier tel quel. Un CSV l'obligeait à reconstruire un
        classeur — la boucle de correction était cassée à l'endroit précis où
        elle devait être la plus fluide.

        La colonne « Motif du rejet » n'est pas reconnue à la relecture : elle
        peut rester, elle sera ignorée. Et comme le dédoublonnage porte sur le
        contenu du bon, renvoyer le fichier ne crée aucun doublon même si une
        ligne avait finalement été reprise entre-temps.
        """
        if not errors or openpyxl is None:
            return False

        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        sheet = workbook.create_sheet("À corriger")

        sheet.cell(1, 1, "Lignes non reprises — corrigez les cellules en rouge, "
                         "puis renvoyez ce fichier à l'assistant d'import.").font = \
            Font(bold=True, color="C0392B")
        header_row = 3

        headers = _TEMPLATE_HEADERS + ["Motif du rejet", "Origine"]
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="305496")
        for column, title in enumerate(headers, start=1):
            cell = sheet.cell(header_row, column, title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for column, width in enumerate(_TEMPLATE_WIDTHS + [42, 28], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = sheet.cell(header_row + 1, 1).coordinate

        culprit_fill = PatternFill("solid", fgColor="F9D6D5")
        culprit_font = Font(bold=True, color="C0392B")
        culprit_column = {'date': 2, 'vehicle': 3, 'beneficiary': 4,
                          'fuel': 5, 'price': 6, 'quantity': 7}

        for index, error in enumerate(errors):
            row_index = header_row + 1 + index
            values = [
                error.get('raw_name'), error.get('raw_date'),
                error.get('vehicle'), error.get('beneficiary'),
                error.get('fuel'), error.get('price'), error.get('quantity'),
                error.get('total'), error.get('error'),
                "%s, ligne %s" % (error.get('sheet') or '', error.get('row') or ''),
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, column, value)
                if column == 2 and isinstance(value, (date, datetime)):
                    cell.number_format = 'DD/MM/YYYY'
            column = culprit_column.get(error.get('culprit'))
            if column:
                faulty = sheet.cell(row_index, column)
                faulty.fill = culprit_fill
                faulty.font = culprit_font
            sheet.cell(row_index, 9).font = Font(color="C0392B")

        stream = io.BytesIO()
        workbook.save(stream)
        return base64.b64encode(stream.getvalue())

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    @api.onchange('file_data')
    def _onchange_file_data(self):
        """Un nouveau fichier remet l'assistant au point de départ."""
        if self.state != 'upload':
            self.state = 'upload'
            self.sheet_ids = [(5, 0, 0)]
