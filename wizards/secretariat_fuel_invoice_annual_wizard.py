# -*- coding: utf-8 -*-
"""État annuel des factures de stations — PDF et Excel.

Un seul tableau répond à la question que pose la direction en fin d'année :
combien a-t-on payé de carburant, mois par mois, station par station ? Le
même jeu de données alimente les deux sorties, pour qu'un chiffre lu dans le
PDF soit exactement celui du classeur.

Le **montant retenu** est explicite : soit le total des bons enregistrés
(ce que le secrétariat a constaté), soit le montant réclamé par la station
(ce qui a été facturé). Les deux diffèrent précisément là où il y a un
problème — d'où la colonne d'écart, toujours présente.
"""

import base64
import io

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError

from ..models.secretariat_dashboard import MONTH_SHORT_FR
from ..models.secretariat_fuel_voucher import MONTH_NAMES_FR

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - dépendance déclarée au manifeste
    openpyxl = None


class SecretariatFuelInvoiceAnnualWizard(models.TransientModel):
    _name = 'secretariat.fuel.invoice.annual.wizard'
    _description = "État annuel des factures de stations"

    year = fields.Integer(
        string="Année", required=True,
        default=lambda self: fields.Date.context_today(self).year)
    station_ids = fields.Many2many(
        'res.partner',
        relation='secretariat_fuel_annual_station_rel',
        column1='wizard_id', column2='partner_id',
        string="Stations",
        domain="[('is_company', '=', True)]",
        help="Laisser vide pour toutes les stations.")
    amount_basis = fields.Selection(
        [
            ('vouchers', "Total des bons enregistrés"),
            ('statement', "Montant réclamé par la station"),
        ],
        string="Montant retenu", default='vouchers', required=True,
        help="Le montant réclamé retombe sur le total des bons quand le "
             "relevé de la station n'a pas été saisi.")
    include_draft = fields.Boolean(
        string="Inclure les factures en brouillon", default=True,
        help="Décochez pour n'arrêter l'état que sur les factures validées "
             "ou payées.")
    compare_previous = fields.Boolean(
        string="Comparer à l'année précédente", default=True)

    state = fields.Selection(
        [('config', "Paramètres"), ('done', "Fichier prêt")],
        default='config', required=True)
    file_data = fields.Binary(string="Classeur", readonly=True, attachment=False)
    file_name = fields.Char(string="Nom du fichier", readonly=True)

    # =========================================================================
    # ACTIONS
    # =========================================================================

    @api.constrains('year')
    def _check_year(self):
        current = fields.Date.context_today(self).year
        for wizard in self:
            if not 2000 <= wizard.year <= current + 1:
                raise ValidationError(_(
                    "L'année %s n'est pas plausible.", wizard.year))

    def action_print(self):
        self.ensure_one()
        if not self._invoices():
            raise UserError(_(
                "Aucune facture de station en %(year)s avec ces filtres.",
                year=self.year))
        return self.env.ref(
            'ivorycocoa_secretariat.action_report_fuel_invoice_annual'
        ).report_action(self, config=False)

    def action_generate_xlsx(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_(
                "La bibliothèque Python « openpyxl » est absente du serveur : "
                "la génération Excel est indisponible."))
        if not self._invoices():
            raise UserError(_(
                "Aucune facture de station en %(year)s avec ces filtres.",
                year=self.year))
        content = self._build_workbook()
        self.write({
            'state': 'done',
            'file_data': base64.b64encode(content),
            'file_name': "factures_stations_%s.xlsx" % self.year,
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    # =========================================================================
    # SÉLECTION
    # =========================================================================

    def _domain(self, year=None):
        self.ensure_one()
        year = year or self.year
        states = ['draft', 'confirmed', 'paid'] if self.include_draft \
            else ['confirmed', 'paid']
        domain = [
            ('date_to', '>=', fields.Date.to_date('%04d-01-01' % year)),
            ('date_to', '<=', fields.Date.to_date('%04d-12-31' % year)),
            ('state', 'in', states),
        ]
        if self.station_ids:
            domain.append(('station_id', 'in', self.station_ids.ids))
        return domain

    def _invoices(self, year=None):
        return self.env['secretariat.fuel.invoice'].search(
            self._domain(year), order='date_to asc, station_id, id asc')

    def _amount_of(self, invoice):
        """Montant retenu pour cette facture, selon la base choisie."""
        self.ensure_one()
        if self.amount_basis == 'statement' and invoice.statement_amount:
            return invoice.statement_amount
        return invoice.amount_total

    # =========================================================================
    # DONNÉES — communes au PDF et à l'Excel
    # =========================================================================

    def report_data(self):
        """Matrice stations × mois, plus les totaux et la comparaison N-1."""
        self.ensure_one()
        invoices = self._invoices()
        stations = {}
        month_totals = [0.0] * 13     # index 1..12
        month_counts = [0] * 13
        grand_total = 0.0
        difference_total = 0.0

        for invoice in invoices:
            month = invoice.date_to.month
            amount = self._amount_of(invoice)
            row = stations.setdefault(invoice.station_id.id, {
                'name': invoice.station_id.display_name or _("(sans station)"),
                'months': [0.0] * 13,
                'counts': [0] * 13,
                'total': 0.0,
                'vouchers': 0,
                'difference': 0.0,
            })
            row['months'][month] += amount
            row['counts'][month] += 1
            row['total'] += amount
            row['vouchers'] += invoice.voucher_count
            row['difference'] += invoice.amount_difference
            month_totals[month] += amount
            month_counts[month] += 1
            grand_total += amount
            difference_total += invoice.amount_difference

        rows = sorted(stations.values(), key=lambda row: -row['total'])

        previous_total = 0.0
        if self.compare_previous:
            previous_total = sum(
                self._amount_of(invoice) for invoice in self._invoices(self.year - 1))

        return {
            'year': self.year,
            'months': [
                {'index': index, 'name': MONTH_NAMES_FR[index],
                 'short': MONTH_SHORT_FR[index],
                 'amount': month_totals[index], 'count': month_counts[index]}
                for index in range(1, 13)
            ],
            'rows': rows,
            'month_totals': month_totals,
            'month_counts': month_counts,
            'grand_total': grand_total,
            'difference_total': difference_total,
            'invoice_count': len(invoices),
            'voucher_count': sum(invoices.mapped('voucher_count')),
            'paid_amount': sum(
                self._amount_of(invoice) for invoice in invoices
                if invoice.state == 'paid'),
            'unpaid_amount': sum(
                self._amount_of(invoice) for invoice in invoices
                if invoice.state != 'paid'),
            'previous_total': previous_total,
            'previous_delta': self._delta(grand_total, previous_total),
            'basis_label': dict(self._fields['amount_basis'].selection)[
                self.amount_basis],
            'invoices': invoices,
            'currency': self.env.company.currency_id,
            'company': self.env.company,
        }

    @staticmethod
    def _delta(current, previous):
        if not previous:
            return None
        return round((current - previous) / previous * 100.0, 1)

    # =========================================================================
    # CLASSEUR EXCEL
    # =========================================================================

    def _build_workbook(self):
        self.ensure_one()
        data = self.report_data()
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        self._write_matrix_sheet(workbook, data)
        self._write_detail_sheet(workbook, data)
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def _write_matrix_sheet(self, workbook, data):
        sheet = workbook.create_sheet("Factures %s" % self.year)
        bold = Font(bold=True)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="305496")
        thin = Side(style='thin', color="BFBFBF")
        border = Border(top=thin, bottom=thin, left=thin, right=thin)

        sheet.cell(1, 1, "État annuel des factures de stations — %s" % self.year) \
            .font = Font(bold=True, size=14)
        sheet.cell(2, 1, "Montant retenu : %s" % data['basis_label'])

        header_row = 4
        headers = ["Station"] + [month['name'] for month in data['months']] \
            + ["Total", "Écart"]
        for col, title in enumerate(headers, start=1):
            cell = sheet.cell(header_row, col, title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        row_index = header_row + 1
        for row in data['rows']:
            sheet.cell(row_index, 1, row['name'])
            for month in range(1, 13):
                cell = sheet.cell(row_index, 1 + month, row['months'][month])
                cell.number_format = '#,##0'
            total_cell = sheet.cell(row_index, 14, row['total'])
            total_cell.number_format = '#,##0'
            total_cell.font = bold
            gap_cell = sheet.cell(row_index, 15, row['difference'])
            gap_cell.number_format = '#,##0'
            row_index += 1

        sheet.cell(row_index, 1, "TOTAL").font = bold
        for month in range(1, 13):
            cell = sheet.cell(row_index, 1 + month, data['month_totals'][month])
            cell.number_format = '#,##0'
            cell.font = bold
        grand = sheet.cell(row_index, 14, data['grand_total'])
        grand.number_format = '#,##0'
        grand.font = bold
        gap = sheet.cell(row_index, 15, data['difference_total'])
        gap.number_format = '#,##0'
        gap.font = bold

        row_index += 2
        sheet.cell(row_index, 1, "Factures").font = bold
        sheet.cell(row_index, 2, data['invoice_count'])
        row_index += 1
        sheet.cell(row_index, 1, "Bons facturés").font = bold
        sheet.cell(row_index, 2, data['voucher_count'])
        row_index += 1
        sheet.cell(row_index, 1, "Payé").font = bold
        sheet.cell(row_index, 2, data['paid_amount']).number_format = '#,##0'
        row_index += 1
        sheet.cell(row_index, 1, "Reste à payer").font = bold
        sheet.cell(row_index, 2, data['unpaid_amount']).number_format = '#,##0'
        if self.compare_previous:
            row_index += 1
            sheet.cell(row_index, 1, "Total %s" % (self.year - 1)).font = bold
            sheet.cell(row_index, 2, data['previous_total']).number_format = '#,##0'

        sheet.column_dimensions['A'].width = 30
        for col in range(2, 16):
            sheet.column_dimensions[get_column_letter(col)].width = 13
        sheet.freeze_panes = 'B5'
        return sheet

    def _write_detail_sheet(self, workbook, data):
        sheet = workbook.create_sheet("Détail")
        bold = Font(bold=True)
        headers = ["Référence", "Station", "Période", "Bons", "Total des bons",
                   "Montant réclamé", "Écart", "État", "Payée le",
                   "N° facture station"]
        for col, title in enumerate(headers, start=1):
            cell = sheet.cell(1, col, title)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")

        states = dict(
            self.env['secretariat.fuel.invoice']._fields['state'].selection)
        row_index = 2
        for invoice in data['invoices']:
            values = [
                invoice.name,
                invoice.station_id.display_name or '',
                invoice.period_label,
                invoice.voucher_count,
                invoice.amount_total,
                invoice.statement_amount,
                invoice.amount_difference,
                states.get(invoice.state, invoice.state),
                invoice.payment_date or '',
                invoice.station_invoice_ref or '',
            ]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, col, value)
                if col in (5, 6, 7):
                    cell.number_format = '#,##0'
                elif col == 9 and value:
                    cell.number_format = 'DD/MM/YYYY'
            row_index += 1

        sheet.cell(row_index, 4, sum(data['invoices'].mapped('voucher_count'))).font = bold
        for col, key in ((5, 'amount_total'), (6, 'statement_amount'),
                         (7, 'amount_difference')):
            cell = sheet.cell(row_index, col, sum(data['invoices'].mapped(key)))
            cell.number_format = '#,##0'
            cell.font = bold

        for col, width in enumerate(
                [16, 28, 24, 8, 16, 16, 12, 12, 12, 20], start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = 'A2'
        return sheet
